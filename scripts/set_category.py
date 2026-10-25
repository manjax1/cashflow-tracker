#!/usr/bin/env python3
"""Change the Category of specific ledger rows (matched by SourceRef) directly in
the Drive ledger. Preview by default; --apply downloads, edits, and uploads.

    python scripts/set_category.py --ref <SourceRef> --category "Rental - Income"
    python scripts/set_category.py --ref <SourceRef> --category "Rental - Income" --apply

Use for one-off fixes to rows already in the ledger. For future rows, add a
keyword rule to spending_rules.json and push it (push_rules_to_drive.py).
"""
import argparse
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "src"))
from dotenv import load_dotenv  # noqa: E402
load_dotenv()

import openpyxl  # noqa: E402
from drive_sync import download_ledger, upload_ledger  # noqa: E402

EXCLUDED_FROM_NET = {"One-Off - Non-Recurring (excluded from Net)", "Investments"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ref", action="append", required=True, help="SourceRef (repeatable)")
    ap.add_argument("--category", required=True)
    ap.add_argument("--apply", action="store_true")
    a = ap.parse_args()

    fid = os.environ["GOOGLE_DRIVE_FILE_ID"]
    tmp = os.path.join(tempfile.gettempdir(), "set_category_ledger.xlsx")
    download_ledger(fid, tmp)
    wb = openpyxl.load_workbook(tmp)
    ws = wb["Transactions"]
    h = [c.value for c in ws[1]]
    ref_c, cat_c, inc_c = h.index("SourceRef"), h.index("Category"), h.index("IncludeInNet")
    d_c, desc_c, amt_c = h.index("Date"), h.index("Description"), h.index("Amount")

    found = set()
    for row in ws.iter_rows(min_row=2):
        ref = str(row[ref_c].value or "")
        if ref in a.ref:
            found.add(ref)
            print(f"  {row[d_c].value}  ${row[amt_c].value}  {str(row[desc_c].value)[:60]}")
            print(f"     {row[cat_c].value!r} -> {a.category!r}")
            if a.apply:
                row[cat_c].value = a.category
                row[inc_c].value = a.category not in EXCLUDED_FROM_NET
    for ref in set(a.ref) - found:
        print(f"  ⚠️  SourceRef not found: {ref}")

    if not a.apply:
        print("\n(preview — re-run with --apply to write to Drive)")
        return
    if not found:
        print("Nothing changed.")
        return
    wb.save(tmp)
    upload_ledger(fid, tmp)
    print(f"\n✅ Updated {len(found)} row(s) on Drive.")


if __name__ == "__main__":
    main()
