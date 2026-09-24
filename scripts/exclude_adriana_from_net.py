#!/usr/bin/env python3
"""One-time migration: flip existing Adriana per-property rows to IncludeInNet=False
in the master ledger, so the spreadsheet's Monthly Summary stops counting them as
income (they're a per-property breakdown; the real income is Adriana's monthly
bank deposit).

The app already excludes them at load time and future imports write them excluded
— this only fixes the ALREADY-imported rows in the stored file. Idempotent.

    python scripts/exclude_adriana_from_net.py            # preview
    python scripts/exclude_adriana_from_net.py --apply    # download → edit → upload
"""
import argparse
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "src"))
from dotenv import load_dotenv  # noqa: E402
load_dotenv(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".env"))

import openpyxl  # noqa: E402
from drive_sync import download_ledger, upload_ledger  # noqa: E402


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true", help="write the change back to Drive")
    args = ap.parse_args()

    fid = os.environ["GOOGLE_DRIVE_FILE_ID"]
    tmp = os.path.join(tempfile.gettempdir(), "adriana_net_migration.xlsx")
    download_ledger(fid, tmp)
    wb = openpyxl.load_workbook(tmp)
    ws = wb["Transactions"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ref_i = header.index("SourceRef")
    inc_i = header.index("IncludeInNet")

    changed = 0
    for row in ws.iter_rows(min_row=2):
        ref = row[ref_i].value
        if ref and str(ref).startswith("adriana:") and row[inc_i].value not in (False, "FALSE", 0):
            if args.apply:
                row[inc_i].value = False
            changed += 1

    print(f"Adriana rows to exclude from net: {changed}")
    if not args.apply:
        print("(dry run — re-run with --apply to write to Drive)")
        return
    wb.save(tmp)
    upload_ledger(fid, tmp)
    print(f"✅ Set IncludeInNet=False on {changed} Adriana rows and uploaded to Drive.")


if __name__ == "__main__":
    main()
