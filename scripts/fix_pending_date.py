#!/usr/bin/env python3
"""Correct the purchase date of a queued (pending) Costco receipt in the master
ledger's PendingReceipts sheet — both the Date column (what the web drawer shows)
AND the date inside ReceiptJSON (what reconciliation actually matches on).

Why this exists: vision extraction can misread the faded transaction date at the
very bottom of a Costco thermal receipt (e.g. the day 04 -> 08). A wrong YEAR or
MONTH makes the queued receipt never match its real charge (reconciliation needs
exact YYYY-MM and day within +/-4), so it would sit until it expires. This pins
the queued row to the true receipt date.

Matches the row by TOTAL (robust to the wrong stored date). Add --receipt-id to
disambiguate if you have more than one pending receipt at the same total.

    # preview (no writes):
    python scripts/fix_pending_date.py --total 222.40 --date 2026-09-04
    # apply (download master -> edit -> upload back to Drive):
    python scripts/fix_pending_date.py --total 222.40 --date 2026-09-04 --apply

Run it from the repo root, in the same environment your daily sync/push uses
(Drive credentials in .env). It is idempotent and safe to re-run.
"""
import argparse
import json
import os
import re
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "src"))

from dotenv import load_dotenv  # noqa: E402

load_dotenv(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".env"))

import openpyxl  # noqa: E402
from drive_sync import download_ledger, upload_ledger  # noqa: E402

HEADERS = ["ReceiptID", "Date", "Total", "Type", "Items", "QueuedAt", "ReceiptJSON"]
ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--total", type=float, required=True,
                    help="grand total of the receipt to fix, e.g. 222.40")
    ap.add_argument("--date", required=True, help="correct purchase date, YYYY-MM-DD")
    ap.add_argument("--receipt-id", default=None,
                    help="optional: exact ReceiptID if multiple rows share the total")
    ap.add_argument("--tol", type=float, default=0.01, help="total match tolerance (default 0.01)")
    ap.add_argument("--apply", action="store_true", help="write the change back to Drive")
    args = ap.parse_args()

    if not ISO.match(args.date):
        sys.exit(f"--date must be YYYY-MM-DD, got {args.date!r}")

    fid = os.environ["GOOGLE_DRIVE_FILE_ID"]
    tmp = os.path.join(tempfile.gettempdir(), "master_fix_pending.xlsx")
    download_ledger(fid, tmp)
    wb = openpyxl.load_workbook(tmp)
    if "PendingReceipts" not in wb.sheetnames:
        sys.exit("No PendingReceipts sheet — nothing queued.")
    ws = wb["PendingReceipts"]
    hdr = [c.value for c in ws[1]]
    idx = {h: hdr.index(h) for h in HEADERS if h in hdr}

    matches = []
    for row in range(2, ws.max_row + 1):
        rid = ws.cell(row=row, column=idx["ReceiptID"] + 1).value
        if rid in (None, ""):
            continue
        total = ws.cell(row=row, column=idx["Total"] + 1).value
        try:
            close = abs(abs(float(total)) - abs(args.total)) <= args.tol
        except (TypeError, ValueError):
            close = False
        if close and (args.receipt_id is None or str(rid) == str(args.receipt_id)):
            matches.append(row)

    if not matches:
        sys.exit(f"No pending receipt found with total ~= {args.total:.2f}"
                 + (f" and id {args.receipt_id}" if args.receipt_id else ""))
    if len(matches) > 1:
        print("Multiple matches — re-run with --receipt-id to pick one:")
        for row in matches:
            print("  id=", ws.cell(row=row, column=idx["ReceiptID"] + 1).value,
                  "date=", ws.cell(row=row, column=idx["Date"] + 1).value,
                  "total=", ws.cell(row=row, column=idx["Total"] + 1).value)
        sys.exit(1)

    row = matches[0]
    old_date = ws.cell(row=row, column=idx["Date"] + 1).value
    rj_cell = ws.cell(row=row, column=idx["ReceiptJSON"] + 1)
    try:
        receipt = json.loads(rj_cell.value)
        old_json_date = receipt.get("date")
    except (TypeError, ValueError, json.JSONDecodeError):
        receipt, old_json_date = None, "<unreadable JSON>"

    print(f"Match: id={ws.cell(row=row, column=idx['ReceiptID']+1).value} "
          f"total={ws.cell(row=row, column=idx['Total']+1).value}")
    print(f"  Date column   : {old_date!r}  ->  {args.date!r}")
    print(f"  ReceiptJSON.date: {old_json_date!r}  ->  {args.date!r}")

    if not args.apply:
        print("\n(dry run — re-run with --apply to write to Drive)")
        return

    ws.cell(row=row, column=idx["Date"] + 1).value = args.date
    if receipt is not None:
        receipt["date"] = args.date
        rj_cell.value = json.dumps(receipt)
    wb.save(tmp)
    upload_ledger(fid, tmp)
    print("\n✅ Applied and uploaded to Drive. It will auto-split at the next sync "
          "once the matching charge posts.")


if __name__ == "__main__":
    main()
