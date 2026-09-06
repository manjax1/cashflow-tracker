"""Durable pending-recategorization proposals.

The agent's action tools (recategorize_transaction / recategorize_batch) are
PROPOSAL-ONLY: nothing is written until a human approves. Historically those
proposals lived only in the immediate /api/chat response, so a proposal made in
one turn vanished the moment the user sent another message — the agent still
"believed" it existed, but the UI and server had discarded it.

This module persists proposals in a 'PendingProposals' sheet in the master
ledger (same pattern as Costco PendingReceipts), so they survive across chat
turns AND Railway redeploys, and can be surfaced in a dedicated web panel until
they're approved or discarded. Keyed by SourceRef — re-proposing the same
transaction replaces its row (idempotent).
"""
import json
from datetime import datetime

from . import ledger
from .tools import audit

_HEADERS = ["ProposalID", "SourceRef", "NewCategory", "Reason", "CreatedAt", "CreatedBy"]


def _iter_items(items):
    for it in items or []:
        ref = str(it.get("source_ref") or it.get("SourceRef") or "").strip()
        cat = (it.get("new_category") or it.get("NewCategory") or "").strip()
        if ref and cat:
            yield ref, cat, (it.get("reason") or "").strip()


def queue_proposals(items, ledger_path, user=""):
    """Upsert proposals (list of {source_ref, new_category, reason}) into the
    PendingProposals sheet, keyed by SourceRef. Returns the number written."""
    import openpyxl
    rows = list(_iter_items(items))
    if not rows:
        return 0
    wb = openpyxl.load_workbook(ledger_path)
    if "PendingProposals" in wb.sheetnames:
        ws = wb["PendingProposals"]
    else:
        ws = wb.create_sheet("PendingProposals")
        ws.append(_HEADERS)
    refs = {r[0] for r in rows}
    for r in range(ws.max_row, 1, -1):                 # drop any prior row for these refs
        if str(ws.cell(row=r, column=2).value) in refs:
            ws.delete_rows(r, 1)
    now = datetime.now().isoformat(timespec="seconds")
    for ref, cat, reason in rows:
        ws.append([f"p{int(datetime.now().timestamp()*1000)}", ref, cat, reason, now, user])
    wb.save(ledger_path)
    ledger._cache["mtime"] = None
    audit("proposals_queued", {"count": len(rows), "refs": sorted(refs), "user": user})
    return len(rows)


def list_proposals(ledger_path):
    """Return queued proposals enriched with each transaction's current
    date/amount/description/category, for the web panel. Self-cleans display:
    a proposal whose transaction is gone, or whose category already equals the
    proposed one (already applied elsewhere), is omitted."""
    import openpyxl
    wb = openpyxl.load_workbook(ledger_path, read_only=True)
    if "PendingProposals" not in wb.sheetnames:
        wb.close()
        return []
    raw = []
    for r in wb["PendingProposals"].iter_rows(min_row=2, values_only=True):
        if r and r[1]:
            raw.append(dict(zip(_HEADERS, r)))
    wb.close()
    by_ref = {str(t["SourceRef"]): t for t in ledger.load_transactions()}
    out = []
    for p in raw:
        ref = str(p["SourceRef"])
        tx = by_ref.get(ref)
        if not tx:
            continue
        if (tx["Category"] or "") == (p["NewCategory"] or ""):
            continue                                   # already in the proposed state
        out.append({"proposal_id": p["ProposalID"], "source_ref": ref,
                    "new_category": p["NewCategory"], "reason": p.get("Reason", ""),
                    "date": tx["Date"], "amount": tx["Amount"],
                    "description": (tx["Description"] or "")[:70],
                    "old_category": tx["Category"], "created_at": p.get("CreatedAt", "")})
    return out


def remove_proposals(ledger_path, source_refs):
    """Delete queued proposals by SourceRef (on apply or discard). Returns count."""
    import openpyxl
    refs = {str(x) for x in source_refs}
    wb = openpyxl.load_workbook(ledger_path)
    if "PendingProposals" not in wb.sheetnames:
        wb.close()
        return 0
    ws = wb["PendingProposals"]
    removed = 0
    for r in range(ws.max_row, 1, -1):
        if str(ws.cell(row=r, column=2).value) in refs:
            ws.delete_rows(r, 1)
            removed += 1
    if removed:
        wb.save(ledger_path)
        ledger._cache["mtime"] = None
        audit("proposals_removed", {"count": removed, "refs": sorted(refs)})
    return removed
