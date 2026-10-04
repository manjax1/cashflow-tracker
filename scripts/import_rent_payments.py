#!/usr/bin/env python3
"""Import authoritative per-unit rental payments (rent_payments_import.json) into
the master ledger, so the Rent view reflects true per-period collections from
Jan 2026. The synced ledger is missing pre-Adriana-itemization detail and buckets
by transaction date rather than rent period; this fixes both.

How it works:
- Each payment is written as a Rental-Income row attributed to the unit's
  account, DATED to the rent PERIOD it's for (so months bucket correctly), with
  SourceRef 'adriana:<period>:<label>:…' and IncludeInNet=False. Income totals
  are unaffected (the consolidated bank deposit is the real income; these rows
  are breakdown/attribution only — no double counting).
- For every unit listed, it first REMOVES existing 'adriana:2026-…' rows on that
  unit's account, then writes the authoritative set. So the file is the source of
  truth; re-running is idempotent.

    python scripts/import_rent_payments.py            # preview (no write)
    python scripts/import_rent_payments.py --apply    # download -> edit -> upload to Drive
"""
import argparse
import json
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "src"))
from dotenv import load_dotenv  # noqa: E402
load_dotenv()

import openpyxl  # noqa: E402
from drive_sync import download_ledger, upload_ledger  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "rent_payments_import.json")
COLS = ["Date", "Description", "Account", "Category", "Type", "Amount", "IncludeInNet", "SourceRef"]


def build_rows(units):
    """Turn the import file into ledger row dicts (period-dated, excluded-from-net)."""
    out = []
    for label, cfg in units.items():
        acct = cfg.get("account_label", label)
        for i, p in enumerate(cfg["payments"], 1):
            ym = str(p["period"])                      # e.g. 2026-04
            day = int(str(p.get("date", ym + "-01"))[8:10] or "1")
            rdate = f"{ym}-{min(max(day, 1), 28):02d}"  # keep day, force month = period
            amt = round(float(p["amount"]), 2)
            out.append({
                "Date": rdate, "Description": (p.get("payer") or "(rent)"),
                "Account": acct, "Category": "Rental - Income", "Type": "Income",
                "Amount": amt, "IncludeInNet": False,
                "SourceRef": f"adriana:{ym}:{label}:{rdate}:{amt:.2f}:{i}",
            })
    return out


def apply_import(ledger_path, units, apply=False):
    rows = build_rows(units)
    wb = openpyxl.load_workbook(ledger_path)
    ws = wb["Transactions"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    ci = {h: header.index(h) for h in COLS}
    accts = {cfg.get("account_label", lbl) for lbl, cfg in units.items()}

    removed = 0
    for r in range(ws.max_row, 1, -1):
        sref = ws.cell(row=r, column=ci["SourceRef"] + 1).value
        acct = ws.cell(row=r, column=ci["Account"] + 1).value
        if sref and str(sref).startswith("adriana:2026-") and acct in accts:
            if apply:
                ws.delete_rows(r, 1)
            removed += 1

    if apply:
        for d in rows:
            vals = [None] * len(header)
            for h in COLS:
                vals[ci[h]] = d[h]
            ws.append(vals)
        wb.save(ledger_path)
    wb.close()
    # period totals for the preview
    totals = {}
    for d in rows:
        totals.setdefault(d["SourceRef"].split(":")[2], {}).setdefault(d["Date"][:7], 0.0)
        totals[d["SourceRef"].split(":")[2]][d["Date"][:7]] += d["Amount"]
    return {"removed": removed, "added": len(rows), "accounts": sorted(accts),
            "period_totals": totals}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true", help="write to Drive")
    args = ap.parse_args()
    units = json.load(open(DATA))["units"]
    fid = os.environ["GOOGLE_DRIVE_FILE_ID"]
    tmp = os.path.join(tempfile.gettempdir(), "rent_import_ledger.xlsx")
    download_ledger(fid, tmp)
    res = apply_import(tmp, units, apply=args.apply)
    print(f"Units: {', '.join(res['accounts'])}")
    print(f"Existing adriana:2026 rows on those accounts {'removed' if args.apply else 'to remove'}: {res['removed']}")
    print(f"Rows {'written' if args.apply else 'to write'}: {res['added']}")
    for label, months in res["period_totals"].items():
        print(f"  {label}: " + ", ".join(f"{m} ${v:,.2f}" for m, v in sorted(months.items())))
    if not args.apply:
        print("\n(preview — re-run with --apply to download, edit, and upload to Drive)")
        return
    upload_ledger(fid, tmp)
    print("\n✅ Applied and uploaded to Drive. The Rent view will reflect it (web: immediately; "
          "spreadsheet tabs: next sync).")


if __name__ == "__main__":
    main()
